#!/usr/bin/env python3
"""Build the eBike QA + store-readiness progress workbook."""

from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, Fill, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart.series import DataPoint
from openpyxl.drawing.fill import PatternFillProperties, ColorChoice
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.marker import DataPoint as DP

OUT = "/Users/arielyakobov/Desktop/ebike/docs/eBike-QA-Progress-2026-08.xlsx"

# Status values (Hebrew, consistent across sheets)
S_CODE = "קיים בקוד — לא נבדק במכשיר"
S_OK = "עובד (ידוע מהיסטוריה / פרודקשן)"
S_PART = "חלקי"
S_BROKEN = "לא עובד / שבור"
S_BLOCK = "בלוקר פרודקשן"
S_MISSING = "חסר לחנות"
S_UX = "UX חלקי / לא אחיד"

STATUSES = [S_CODE, S_OK, S_PART, S_BROKEN, S_BLOCK, S_MISSING, S_UX]

PRI = ["P0 — חוסם השקה", "P1 — חשוב", "P2 — שיפור"]

thin = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)
wrap = Alignment(wrap_text=True, vertical="top", readingOrder=2)
wrap_ltr = Alignment(wrap_text=True, vertical="top")
center = Alignment(wrap_text=True, vertical="center", horizontal="center", readingOrder=2)

fills = {
    "header": PatternFill("solid", fgColor="1E3A8A"),
    "header2": PatternFill("solid", fgColor="1D4ED8"),
    "p0": PatternFill("solid", fgColor="FEE2E2"),
    "p1": PatternFill("solid", fgColor="FEF3C7"),
    "p2": PatternFill("solid", fgColor="DBEAFE"),
    "ok": PatternFill("solid", fgColor="D1FAE5"),
    "part": PatternFill("solid", fgColor="FEF3C7"),
    "broken": PatternFill("solid", fgColor="FECACA"),
    "block": PatternFill("solid", fgColor="FCA5A5"),
    "missing": PatternFill("solid", fgColor="E9D5FF"),
    "code": PatternFill("solid", fgColor="E0E7FF"),
    "ux": PatternFill("solid", fgColor="CFFAFE"),
    "alt": PatternFill("solid", fgColor="F8FAFC"),
    "note": PatternFill("solid", fgColor="EFF6FF"),
    "title": PatternFill("solid", fgColor="0F172A"),
}

font_h = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
font_title = Font(name="Calibri", bold=True, color="FFFFFF", size=18)
font_sub = Font(name="Calibri", bold=True, color="1E3A8A", size=13)
font_n = Font(name="Calibri", size=11)
font_b = Font(name="Calibri", bold=True, size=11)


def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = fills["header"]
        cell.font = font_h
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = thin


def style_body(ws, start, end, cols, status_col=None, pri_col=None):
    for r in range(start, end + 1):
        for c in range(1, cols + 1):
            cell = ws.cell(r, c)
            cell.font = font_n
            cell.alignment = wrap
            cell.border = thin
            if r % 2 == 0:
                cell.fill = fills["alt"]
        if status_col:
            v = ws.cell(r, status_col).value or ""
            fill = None
            if v == S_OK:
                fill = fills["ok"]
            elif v == S_PART:
                fill = fills["part"]
            elif v == S_BROKEN:
                fill = fills["broken"]
            elif v == S_BLOCK:
                fill = fills["block"]
            elif v == S_MISSING:
                fill = fills["missing"]
            elif v == S_CODE:
                fill = fills["code"]
            elif v == S_UX:
                fill = fills["ux"]
            if fill:
                ws.cell(r, status_col).fill = fill
                ws.cell(r, status_col).font = font_b
        if pri_col:
            v = ws.cell(r, pri_col).value or ""
            if v.startswith("P0"):
                ws.cell(r, pri_col).fill = fills["p0"]
            elif v.startswith("P1"):
                ws.cell(r, pri_col).fill = fills["p1"]
            elif v.startswith("P2"):
                ws.cell(r, pri_col).fill = fills["p2"]


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_status_dv(ws, col, start, end):
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    dv.error = "בחר סטטוס מהרשימה"
    dv.errorTitle = "סטטוס"
    ws.add_data_validation(dv)
    dv.add(f"{col}{start}:{col}{end}")


def add_pri_dv(ws, col, start, end):
    dv = DataValidation(type="list", formula1='"' + ",".join(PRI) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}{start}:{col}{end}")


def add_yn_dv(ws, col, start, end):
    dv = DataValidation(type="list", formula1='"כן,לא,חלקי,לא רלוונטי"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}{start}:{col}{end}")


wb = Workbook()

# ═══════════════════════════════════════════════════════════════════════════
# 00 README
# ═══════════════════════════════════════════════════════════════════════════
ws = wb.active
ws.title = "00_קרא_אותי"
ws.sheet_view.rightToLeft = True
ws.merge_cells("A1:F1")
ws["A1"] = "eBike — דוח QA + התקדמות השקה  |  18 באוגוסט 2026"
ws["A1"].font = font_title
ws["A1"].fill = fills["title"]
ws["A1"].alignment = Alignment(vertical="center", readingOrder=2)
ws.row_dimensions[1].height = 32

ws.merge_cells("A2:F2")
ws["A2"] = (
    "מקור: בדיקת קוד מלאה + /health של פרודקשן + היסטוריית סשנים. "
    "זה לא E2E חי על שני מכשירים — העמודה «נבדק במכשיר» ריקה בכוונה, כדי שתמלא אחרי שתפתח את האפליקציה."
)
ws["A2"].fill = fills["note"]
ws["A2"].alignment = wrap
ws.row_dimensions[2].height = 42

readme_rows = [
    ("מה הסטטוסים אומרים", ""),
    (S_CODE, "הפיצ'ר קיים בקוד. לא הרצתי אותו עכשיו על סימולטור/TestFlight."),
    (S_OK, "עבד בעבר בסימולטור / ידוע כעובד, או מאושר מ־/health."),
    (S_PART, "עובד חלקית — חסר סניף, באג ידוע, או תלוי סביבה."),
    (S_BROKEN, "שבור או נכשל בהיסטוריה ולא תוקן בפרודקשן."),
    (S_BLOCK, "חוסם השקה לחנות / תשלום אמיתי."),
    (S_MISSING, "חסר ל־App Store / Play (מסמכים, חשבון, נכסים)."),
    (S_UX, "עובד פונקציונלית אבל העיצוב לא אחיד / לא מלוטש."),
    ("", ""),
    ("איך להתקדם עם הקובץ", ""),
    ("1", "פתח «01_איפה_אנחנו» — תמונת מצב אחרי חודש בלי האפליקציה."),
    ("2", "עבור על «07_בלוקרים» — אלה הדברים שחייבים לפני חנות."),
    ("3", "בזמן בדיקה במכשיר: ב«02_מאסטר_QA» עדכן «נבדק במכשיר» ל־כן/לא ואת הסטטוס."),
    ("4", "«05_אפל_100» ו«06_גוגל_100» — רשימת השקה מלאה. סמן V כשסיימת."),
    ("5", "גיליון «סיכום» מתעדכן אוטומטית לפי הסטטוסים בגיליון המאסטר."),
    ("", ""),
    ("חשבונות בדיקה", ""),
    ("לקוח", "a@abc.com / 123456"),
    ("טכנאי", "tech-test@ebike.com / 123456"),
    ("אלביס / מאור", "maort@ebike.com"),
    ("Backend פרודקשן", "https://ebikel-backend.onrender.com"),
    ("/health עכשיו", 'status=ok | version=2026-07-28-payment-500-fix | mockPayments=TRUE | provider=mock'),
    ("iOS bundle", "com.ebikeland.app  |  ASC 6775995887  |  גרסה 1.8.1 build 24"),
    ("Android package", "com.ebike.app  |  versionCode 24  |  אין Play Console עדיין"),
]

ws["A4"] = "נושא"
ws["B4"] = "פירוט"
style_header(ws, 4, 2)
for i, (a, b) in enumerate(readme_rows, 5):
    ws.cell(i, 1, a)
    ws.cell(i, 2, b)
    for c in range(1, 3):
        ws.cell(i, c).alignment = wrap
        ws.cell(i, c).border = thin
        ws.cell(i, c).font = font_n
    if a in fills or a.startswith("קיים") or a.startswith("עובד") or a.startswith("חלקי") or a.startswith("לא ") or a.startswith("בלוקר") or a.startswith("חסר") or a.startswith("UX"):
        ws.cell(i, 1).fill = (
            fills["ok"] if a == S_OK else
            fills["part"] if a == S_PART else
            fills["broken"] if a == S_BROKEN else
            fills["block"] if a == S_BLOCK else
            fills["missing"] if a == S_MISSING else
            fills["code"] if a == S_CODE else
            fills["ux"] if a == S_UX else
            fills["alt"]
        )
        ws.cell(i, 1).font = font_b
    if a in ("מה הסטטוסים אומרים", "איך להתקדם עם הקובץ", "חשבונות בדיקה"):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
        ws.cell(i, 1).fill = fills["header2"]
        ws.cell(i, 1).font = font_h

autosize(ws, [36, 92])
ws.freeze_panes = "A5"
ws.row_dimensions[4].height = 22

# ═══════════════════════════════════════════════════════════════════════════
# 01 WHERE WE ARE
# ═══════════════════════════════════════════════════════════════════════════
ws = wb.create_sheet("01_איפה_אנחנו")
ws.sheet_view.rightToLeft = True
ws.merge_cells("A1:D1")
ws["A1"] = "איפה היינו כשעזבת (~חודש) ואיפה אנחנו היום"
ws["A1"].font = font_title
ws["A1"].fill = fills["title"]
ws.row_dimensions[1].height = 28

headers = ["נושא", "מצב אחרון שידעת", "מצב היום (18.08.2026)", "מה זה אומר לך"]
for i, h in enumerate(headers, 1):
    ws.cell(2, i, h)
style_header(ws, 2, 4)

where = [
    ("גרסת iOS בחנות / TF", "1.8.0 build 21–23, ASC חדש 6775995887", "app.json = 1.8.1 build 24. EAS build+submit הורץ (6931d93c…). לבדוק ב-TestFlight אם הגיע.", "פתח TestFlight. אם 1.8.1 שם — זה הבילד עם UI זכוכית + תיקון נוסף."),
    ("Backend Render", "PayMe אמור להיות מחובר; לעיתים mock", "GET /health → mockPayments=true, provider=mock, version=2026-07-28-payment-500-fix", "בחנות אמיתית זה בלוקר: הלקוחות ישלמו בדף דמו, לא בכרטיס."),
    ("תשלומים", "PayMe sandbox עבד מקומית; פרודקשן התנדנד", "קוד PayMe קיים. פרודקשן רץ mock. extra-repair נוסף בקוד.", "להגדיר MOCK_PAYMENTS=false + מפתחות PayMe ב-Render ולעשות deploy."),
    ("זרימת הזמנה", "לקוח משלם רק אחרי accept, לפני on_way", "עדיין הכלל. שערים בשרת + מסכי Waiting / Pay / Tracking", "הלוגיקה נכונה. חובה E2E אחרי PayMe אמיתי."),
    ("UI", "מסך בית + התחברות יפים; שאר המסכים גנריים", "זכוכית: splash, בית, תיקון, בחירת טכנאי, תשלום, השלמה, מודאלים. לא כל האפליקציה אחידה.", "פרופיל / הזמנות / רווחים / אדמין עדיין ישנים."),
    ("טכנאי נעלם מהרשימה", "תוקן בקוד (sync + רדיוס 40)", "הקוד המקומי כולל את התיקון. פרודקשן backend ישן יחסית (יולי).", "לוודא שה-deploy ב-Render כולל רדיוס/sync."),
    ("חשבון טכנאי", "נרשמים כלקוח; מנהל מאשר טכנאי", "כך גם ב-legal. אדמין בטאב הטכנאי.", "אלביס/מאור צריכים אישור אדמין."),
    ("Google Play", "לא התחלנו", "אין אפליקציה בקונסול. package=com.ebike.app", "קודם חשבון Play $25 ואז create app."),
    ("App Store", "ITMS-90055 → אפליקציה חדשה 6775995887", "iOS מוכן טכנית. חסר: הסכם Apple אם צריך, צילומים, סקירה, Data.", "TestFlight → Internal → Submit for Review."),
    ("קוד לא commited", "הרבה שינויי UI מקומיים", "git status היה dirty בשיחה הקודמת (UI, extra pay, splash)", "לפני ביילד נוסף: commit + לוודא ש-EAS לקח את הקבצים."),
]
for i, row in enumerate(where, 3):
    for c, v in enumerate(row, 1):
        ws.cell(i, c, v)
style_body(ws, 3, 2 + len(where), 4)
autosize(ws, [22, 42, 55, 48])
ws.freeze_panes = "A3"
for r in range(3, 3 + len(where)):
    ws.row_dimensions[r].height = 48

# ═══════════════════════════════════════════════════════════════════════════
# 02 MASTER QA
# ═══════════════════════════════════════════════════════════════════════════
ws = wb.create_sheet("02_מאסטר_QA")
ws.sheet_view.rightToLeft = True
ws.merge_cells("A1:J1")
ws["A1"] = "מאסטר QA — כל פעולה קטנה באפליקציה  |  עדכן «נבדק במכשיר» אחרי שתפתח"
ws["A1"].font = font_title
ws["A1"].fill = fills["title"]
ws.row_dimensions[1].height = 28

mh = ["ID", "תחום", "מה בודקים", "סטטוס", "עדיפות", "מה עובד / לא עובד — הסבר", "UX/UI", "איך לבדוק", "נבדק במכשיר", "הערות שלך"]
for i, h in enumerate(mh, 1):
    ws.cell(2, i, h)
style_header(ws, 2, 10)

# (id, area, item, status, pri, explain, ux, how)
qa = [
    ("A01", "פתיחה", "Splash נייטיבי (לפני JS)", S_CODE, "P2 — שיפור", "app.json splash = תמונת ההתחברות cover. נכנס רק אחרי rebuild native (EAS). Fast Refresh לא מעדכן splash.", "רציף עם LoadingScreen אחרי JS", "להרוג אפליקציה לגמרי ולפתוח מחדש מ-TestFlight", ""),
    ("A02", "פתיחה", "LoadingScreen בזמן session/RTL", S_CODE, "P1 — חשוב", "רקע sign-in-bg + לוגו + פס טוען. מוצג מ-_layout כשאין rtlReady או session loading.", "יפה, תואם התחברות", "כניסה קרה / רשת איטית", ""),
    ("A03", "פתיחה", "Welcome / tutorial בפעם הראשונה", S_CODE, "P2 — שיפור", "index מפנה ל-/welcome אם לא נראה. אחרת sign-in.", "מסך welcome כהה — פחות זכוכית משאר האפליקציה", "התקנה נקייה", ""),
    ("A04", "התחברות", "כניסת אימייל+סיסמה", S_OK, "P0 — חוסם השקה", "sign-in + better-auth. /health emailSignUpEnabled=true. חשבונות בדיקה עבדו בעבר. באג ידוע: לפעמים session לא נטען אחרי הצלחה.", "טופס משני מוסתר מאחורי לינק", "a@abc.com / 123456", ""),
    ("A05", "התחברות", "כניסת Google", S_PART, "P0 — חוסם השקה", "כפתור קיים. תלוי GOOGLE_CLIENT_* ו-OAUTH_BASE_URL ב-Render. בהיסטוריה: session null אחרי redirect בסימולטור.", "כפתור 3D יפה", "Google אמיתי במכשיר אמיתי (לא רק סימולטור)", ""),
    ("A06", "התחברות", "כניסת Apple", S_PART, "P0 — חוסם השקה", "חובה ל-App Store אם יש login חברתי. תלוי APPLE_* ב-Render. לא אומת במכשיר בחודש האחרון.", "כפתור שחור תקין", "Sign in with Apple ב-iPhone", ""),
    ("A07", "התחברות", "הרשמה אימייל (sign-up)", S_CODE, "P1 — חשוב", "מסך sign-up קיים. כל משתמש נרשם כלקוח. טכנאי = אישור אדמין.", "רקע זהה להתחברות", "משתמש חדש", ""),
    ("A08", "התחברות", "שגיאות סיסמה שגויה", S_CODE, "P1 — חשוב", "ConfirmModal עם הודעת שרת או «אימייל או סיסמא שגויים».", "מודאל זכוכית חדש", "סיסמה שגויה", ""),
    ("A09", "התחברות", "שמירת session אחרי סגירת אפליקציה", S_PART, "P0 — חוסם השקה", "use-session + placeholderData תוקנו מקומית (false logout). TestFlight 1.8.0 הישן אולי בלי התיקון. 1.8.1 אמור לכלול.", "LoadingScreen במקום מסך שחור", "סגור אפליקציה, פתח שוב", ""),
    ("A10", "התחברות", "התנתקות", S_CODE, "P1 — חשוב", "פרופיל לקוח/טכנאי → ConfirmModal → signOut → sign-in.", "", "התנתק והתחבר שוב", ""),
    ("A11", "התחברות", "מחיקת חשבון", S_CODE, "P0 — חוסם השקה", "DELETE /api/users/me מוחק User. חובה ל-Apple. לבדוק cascade על jobs/payments.", "מודאל אזהרה קיים", "חשבון חדש זמני — מחק", ""),
    ("A12", "ניתוב תפקיד", "לקוח נכנס לטאבים של לקוח", S_CODE, "P0 — חוסם השקה", "index: role technician → tech tabs; אחרת customer. אין role-select לציבור — טכנאי דרך אדמין.", "", "התחבר כלקוח", ""),
    ("A13", "ניתוב תפקיד", "טכנאי מאושר → דשבורד", S_CODE, "P0 — חוסם השקה", "isApproved או isAdmin. אחרת technician-pending.", "", "tech-test@ebike.com", ""),
    ("A14", "ניתוב תפקיד", "טכנאי ממתין לאישור", S_CODE, "P1 — חשוב", "מסך technician-pending + התנתקות.", "עיצוב ישן (צהוב) לא זכוכית", "חשבון טכנאי לא מאושר", ""),

    ("C01", "בית לקוח", "מפה + מיקום נוכחי", S_PART, "P0 — חוסם השקה", "WebView Google Maps. דורש הרשאת מיקום + מפתח Maps. מחוץ לישראל ב-DEV מחליף לברירת מחדל.", "מפה מלאה + header זכוכית צף", "אשר מיקום, ודא נקודה כחולה", ""),
    ("C02", "בית לקוח", "כפתור בקש תיקון עכשיו", S_CODE, "P0 — חוסם השקה", "CTA זכוכית כחולה גדולה → /repair-request. אם מיקום denied — מודאל הגדרות.", "חזק ויזואלית, פונט 22", "לחץ CTA", ""),
    ("C03", "בית לקוח", "כרטיס הזמנה פעילה (Hero)", S_CODE, "P0 — חוסם השקה", "כשיש active job מוסתר CTA ומוצג כרטיס גדול → job-tracking.", "עדיין כרטיס לבן/כחול, פחות זכוכית מה-CTA", "עם הזמנה פתוחה", ""),
    ("C04", "בית לקוח", "הדרכה (tutorial) בפינה", S_CODE, "P2 — שיפור", "הועבר לכותרת ימין ליד רענון.", "קטן אבל במקום נכון", "לחץ הדרכה", ""),
    ("C05", "בית לקוח", "רענון מיקום + טכנאים", S_CODE, "P1 — חשוב", "כפתור refresh בהדר זכוכית.", "", "לחץ רענון", ""),
    ("C06", "בית לקוח", "מרכוז מפה על המשתמש", S_CODE, "P2 — שיפור", "כפתור MapPin מתחת להדר.", "", "", ""),
    ("C07", "בית לקוח", "כרטיס טכנאי שנבחר על המפה", S_CODE, "P2 — שיפור", "לחיצה על מרקר → כרטיס תחתון. תלוי WebView postMessage.", "עיצוב ישן יותר מההדר", "אם מופיעים מרקרים", ""),

    ("R01", "דיווח תקלה", "שלב 1 צילום / גלריה / דלג", S_CODE, "P0 — חוסם השקה", "צילום אופציונלי (אפשר לדלג). העלאה בפועל בזמן הזמנה.", "כפתורי זכוכית", "צלם / דלג", ""),
    ("R02", "דיווח תקלה", "שלב 2 סוג אופניים + קטגוריות", S_CODE, "P0 — חוסם השקה", "רגיל/חשמלי + multi-select. מחיר לפי PRICE_RANGES.", "chips זכוכית", "בחר 2 קטגוריות", ""),
    ("R03", "דיווח תקלה", "שלב 3 שם (חובה)", S_CODE, "P0 — חוסם השקה", "ולידציה + שמירת ברירת מחדל.", "שדות אפורים — פחות זכוכית", "השאר ריק → שגיאה", ""),
    ("R04", "דיווח תקלה", "שלב 3 טלפון (חובה)", S_CODE, "P0 — חוסם השקה", "ניקוי תווים / פורמט ישראלי.", "", "050…", ""),
    ("R05", "דיווח תקלה", "שלב 3 עיר (picker + חיפוש)", S_CODE, "P0 — חוסם השקה", "רשימת ערים קשיחה בישראל + modal חיפוש.", "מודאל עיר לבן ישן", "בחר תל אביב", ""),
    ("R06", "דיווח תקלה", "שלב 3 רחוב (API / ידני)", S_PART, "P1 — חשוב", "טוען מ-/api/streets. אם ריק — הזנה ידנית. תלוי backend.", "מודאל רחוב ישן", "בחר רחוב", ""),
    ("R07", "דיווח תקלה", "שלב 3 מספר בית", S_CODE, "P0 — חוסם השקה", "חובה, מוגבל אורך.", "", "", ""),
    ("R08", "דיווח תקלה", "שלב 3 אימייל אופציונלי", S_CODE, "P2 — שיפור", "ולידציית פורמט אם הוזן.", "", "", ""),
    ("R09", "דיווח תקלה", "כתובות שמורות", S_CODE, "P2 — שיפור", "GET/POST saved addresses. בחירה ממלאת שדות.", "", "שמור כתובת ואז בחר", ""),
    ("R10", "דיווח תקלה", "שמור פרטים לפעם הבאה", S_CODE, "P2 — שיפור", "AsyncStorage + API defaults.", "", "", ""),
    ("R11", "דיווח תקלה", "שלב 4 סיכום מחיר + חפש טכנאי", S_CODE, "P0 — חוסם השקה", "הערכת מחיר. כפתור CTA זכוכית. גיאוקוד כתובת לפני המשך.", "כרטיס סיכום זכוכית", "המשך לבחירת טכנאי", ""),
    ("R12", "דיווח תקלה", "פס התקדמות + נקודות שלבים", S_CODE, "P2 — שיפור", "זכוכית, בלי אחוזים ובלי «שלב X מתוך 4» (הוסרו לפי בקשתך).", "נקי", "", ""),
    ("R13", "דיווח תקלה", "חזרה בין שלבים / back", S_CODE, "P1 — חשוב", "handleBack.", "", "חזור אחורה", ""),

    ("T01", "בחירת טכנאי", "טעינת טכנאים זמינים", S_PART, "P0 — חוסם השקה", "GET available לפי מיקום. רדיוס מינימום 40 בקוד backend — פרודקשן אולי ישן. טכנאי חייב approved+available+מיקום טרי.", "מסך עבר זכוכית", "טכנאי זמין ליד הכתובת", ""),
    ("T02", "בחירת טכנאי", "מפה מיני + מרקרים", S_CODE, "P1 — חשוב", "react-native-maps, קואורדינטות מסוננות (תיקון קראש).", "מסגרת מעוגלת", "", ""),
    ("T03", "בחירת טכנאי", "מיון: קרוב / דירוג / מחיר", S_CODE, "P2 — שיפור", "chips זכוכית.", "", "החלף מיון", ""),
    ("T04", "בחירת טכנאי", "כרטיס טכנאי + בחר והזמן", S_CODE, "P0 — חוסם השקה", "כפתור pill מעוגל, טקסט ממורכז.", "טוב", "לחץ בחר והזמן", ""),
    ("T05", "בחירת טכנאי", "גיליון אישור הזמנה", S_CODE, "P0 — חוסם השקה", "זכוכית כחולה + CTA. יוצר job pending + העלאת תמונה.", "", "אשר הזמנה", ""),
    ("T06", "בחירת טכנאי", "409 — כבר יש הזמנה פעילה", S_CODE, "P0 — חוסם השקה", "מודאל ומעביר ל-job-tracking של הקיימת.", "", "נסה להזמין פעמיים", ""),
    ("T07", "בחירת טכנאי", "אין טכנאים + וואטסאפ נציג", S_CODE, "P1 — חשוב", "שולח פרטים ל-WhatsApp תמיכה.", "", "כבה טכנאים / רדיוס רחוק", ""),
    ("T08", "בחירת טכנאי", "פרופיל טכנאי (modal)", S_CODE, "P2 — שיפור", "technician-profile.tsx.", "לא עבר זכוכית", "לחץ על כרטיס", ""),

    ("J01", "מעקב לקוח", "מסך ממתין לטכנאי (pending)", S_CODE, "P0 — חוסם השקה", "WaitingScreen: «יוצרים קשר… עד 5 דקות / כמה שניות». ביטול גדול.", "זכוכית כהה", "אחרי הזמנה לפני accept", ""),
    ("J02", "מעקב לקוח", "ביטול הזמנה לפני תשלום", S_CODE, "P0 — חוסם השקה", "מותר. אחרי תשלום — חסום באפליקציה, וואטסאפ תמיכה.", "מודאל זכוכית", "בטל ב-pending", ""),
    ("J03", "מעקב לקוח", "הטכנאי מוכן לצאת + שלם עכשיו", S_CODE, "P0 — חוסם השקה", "PaymentRequiredScreen. כפתור שלם ~140px. ביטול מתחת.", "זכוכית כהה + CTA ענק", "אחרי שהטכנאי מאשר", ""),
    ("J04", "מעקב לקוח", "יצירת דף תשלום", S_BLOCK, "P0 — חוסם השקה", "פרודקשן mockPayments=true → דף דמו, לא PayMe. הקוד תומך PayMe.", "WebView", "שלם עכשיו מול Render הנוכחי = mock", ""),
    ("J05", "מעקב לקוח", "הצלחת תשלום — טקסטים", S_CODE, "P0 — חוסם השקה", "«תשלום בוצע בהצלחה / טכנאי בדרך». כפתור סיום זכוכית ירוקה.", "", "סיים תשלום", ""),
    ("J06", "מעקב לקוח", "כישלון תשלום + סיבה", S_CODE, "P0 — חוסם השקה", "«התשלום לא הצליח» + reason מ-URL / ביטול / שגיאת רשת.", "", "בטל באמצע / נתק רשת", ""),
    ("J07", "מעקב לקוח", "מפה + סטטוס + ETA אחרי תשלום", S_CODE, "P0 — חוסם השקה", "poll ~4 שנ'. כפתורי תמיכה/ביטול זכוכית גדולים (ביטול רק אם לא שולם).", "תחתון לבן + כפתורי זכוכית", "אחרי תשלום", ""),
    ("J08", "מעקב לקוח", "שיחה / וואטסאפ לטכנאי", S_CODE, "P1 — חשוב", "ריבועים עם תווית. דורש טלפון בחשבון הטכנאי.", "", "", ""),
    ("J09", "מעקב לקוח", "אישור הגעת טכנאי", S_CODE, "P1 — חשוב", "כשהסטטוס arrived.", "", "", ""),
    ("J10", "מעקב לקוח", "תיקון נוסף — פופאפ + תשלום", S_CODE, "P1 — חשוב", "poll extra-repair pending. אשר ושלם / דחה. חדש — לא נבדק E2E.", "מודאל זכוכית יפה", "טכנאי שולח בקשת תוספת", ""),
    ("J11", "מעקב לקוח", "סיום עבודה + דירוג", S_CODE, "P0 — חוסם השקה", "job-complete זכוכית. כוכבים + קטגוריות + משוב. אפשר לדלג.", "", "טכנאי משלים", ""),
    ("J12", "מעקב לקוח", "צ'אט בהזמנה", S_CODE, "P2 — שיפור", "chat.tsx + messages API. לא חלק מהזרימה הראשית.", "עיצוב כהה ישן", "פתח צ'אט אם יש כניסה", ""),

    ("K01", "טכנאי דשבורד", "מתג זמין/לא זמין", S_PART, "P0 — חוסם השקה", "שולח מיקום + heartbeat. reassertAvailability אחרי רקע. באג ישן: נעלם אחרי נסיעה — תוקן בקוד.", "דשבורד לא זכוכית", "הדלק זמין, ודא שמופיע אצל לקוח", ""),
    ("K02", "טכנאי דשבורד", "רשימת הזמנות ממתינות + קבל", S_CODE, "P0 — חוסם השקה", "poll + באנר הזמנה חדשה + accept.", "", "לקוח מזמין את הטכנאי", ""),
    ("K03", "טכנאי דשבורד", "סטטיסטיקות יום/שבוע", S_CODE, "P2 — שיפור", "API earnings/stats.", "", "", ""),
    ("K04", "טכנאי עבודה", "פרטי לקוח + ניווט סטטוסים", S_CODE, "P0 — חוסם השקה", "on_way רק אחרי paid. arrived → in_progress → complete.", "מסך כהה מלוטש", "עבור שלבים אחרי תשלום", ""),
    ("K05", "טכנאי עבודה", "תיאור תקלה מהלקוח", S_CODE, "P1 — חשוב", "נוסף כרטיס מלא (לא קטוע). כולל כתובת אם קיימת.", "", "הזמנה עם תיאור חופשי", ""),
    ("K06", "טכנאי עבודה", "תמונה מהלקוח", S_CODE, "P1 — חשוב", "מוצגת אם הועלתה.", "", "", ""),
    ("K07", "טכנאי עבודה", "טלפון / וואטסאפ ללקוח", S_CODE, "P1 — חשוב", "רק אחרי paid + on_way/arrived/in_progress/completed.", "", "", ""),
    ("K08", "טכנאי עבודה", "הזמן טכנאי נוסף", S_CODE, "P2 — שיפור", "הזמנות + מסך invitations.", "", "", ""),
    ("K09", "טכנאי עבודה", "בקש תשלום נוסף", S_CODE, "P1 — חשוב", "טופס תיאור+סכום → POST extra-repair. חדש, לא E2E.", "", "arrived/in_progress + paid", ""),
    ("K10", "טכנאי עבודה", "תקלה לא תוקנה — החזר", S_CODE, "P1 — חשוב", "POST refund. מסמן refund_requested.", "", "", ""),
    ("K11", "טכנאי עבודה", "ביטול הזמנה (טכנאי)", S_CODE, "P0 — חוסם השקה", "אחרי תשלום אמור לפתוח תהליך החזר.", "", "", ""),
    ("K12", "טכנאי עבודה", "סיום עבודה + מחיר סופי + חלקים", S_PART, "P0 — חוסם השקה", "טופס complete. באג ישן: «לא הצלחנו לטעון» בזמן סיום — תוקן בקוד (placeholder + poll 4s).", "", "סיים עבודה", ""),
    ("K13", "טכנאי עבודה", "הוסר: רענן סטטוס תשלום", S_OK, "P2 — שיפור", "הוסר לפי בקשתך. הסטטוס מתעדכן ב-poll.", "", "", ""),
    ("K14", "טכנאי", "טאב הזמנות / רווחים / משיכה", S_CODE, "P1 — חשוב", "jobs + earnings + withdrawal-request (פרטי בנק).", "עיצוב ישן", "מלא משיכה בסביבת בדיקה", ""),
    ("K15", "טכנאי", "פרופיל: ביו, רדיוס, מחיר בסיס, זמינות", S_CODE, "P1 — חשוב", "שדות + שמירה.", "לא זכוכית", "", ""),
    ("K16", "אדמין", "אישור / ביטול / הסרת טכנאי", S_CODE, "P0 — חוסם השקה", "טאב admin ל-isAdmin. כך נכנסים טכנאים חדשים.", "", "התחבר כמנהל", ""),

    ("P01", "פרופיל לקוח", "עריכת שם / טלפון", S_CODE, "P1 — חשוב", "edit-profile. אימייל לא ניתן לעריכה.", "שדות פשוטים", "", ""),
    ("P02", "פרופיל לקוח", "כתובות שמורות CRUD", S_CODE, "P2 — שיפור", "saved-addresses.", "", "", ""),
    ("P03", "פרופיל לקוח", "שפה עברית/אנגלית + RTL", S_CODE, "P1 — חשוב", "applyRtlForLanguage. דורש reload.", "חלק מהמסכים לא מושלמים ב-LTR", "החלף שפה", ""),
    ("P04", "פרופיל לקוח", "התראות on/off", S_CODE, "P1 — חשוב", "push preference.", "", "", ""),
    ("P05", "פרופיל לקוח", "מצב כהה", S_PART, "P2 — שיפור", "toggle קיים. הרבה מסכים hard-coded כהים/לבנים — לא כל האפליקציה מגיבה.", "לא עקבי", "החלף ערכת נושא", ""),
    ("P06", "פרופיל לקוח", "תמיכה וואטסאפ", S_CODE, "P0 — חוסם השקה", "+972585858586 דרך support.ts.", "", "פתח וואטסאפ", ""),
    ("P07", "פרופיל לקוח", "תנאים / פרטיות / ביטול", S_PART, "P0 — חוסם השקה", "מסמכים בתוך האפליקציה. חסר URL ציבורי לחנויות.", "מסך legal פשוט", "", ""),
    ("P08", "פרופיל לקוח", "הזמנות קודמות", S_CODE, "P1 — חשוב", "טאב orders + order-details.", "לא זכוכית", "", ""),
    ("P09", "פרופיל לקוח", "העלאת תמונת פרופיל", S_CODE, "P2 — שיפור", "ImagePicker + upload.", "", "", ""),

    ("X01", "תשתית", "Backend /health חי", S_OK, "P0 — חוסם השקה", "נבדק עכשיו: ok. אבל mock=true.", "", "curl /health", ""),
    ("X02", "תשתית", "PayMe בפרודקשן", S_BLOCK, "P0 — חוסם השקה", "קוד קיים. Render רץ mock. בלי זה אין סליקה אמיתית.", "", "שנה env + deploy + בדוק /health", ""),
    ("X03", "תשתית", "העלאת תמונות (Supabase)", S_PART, "P1 — חשוב", "bucket job-photos. תלוי env ב-Render.", "", "העלה תמונה בהזמנה", ""),
    ("X04", "תשתית", "פוש נוטיפיקציות", S_PART, "P0 — חוסם השקה", "Expo push + ניווט. דורש טוקן והרשאה. קריטי לקבלת הזמנה.", "", "שני מכשירים", ""),
    ("X05", "תשתית", "מיקום רקע לטכנאי", S_PART, "P0 — חוסם השקה", "expo-location background + heartbeat. רגיש ל-iOS.", "הסבר הרשאה קיים", "זמין + נעל אפליקציה + לקוח מחפש", ""),
    ("X06", "תשתית", "מסך modal.tsx סטאב", S_UX, "P2 — שיפור", "מסך תבנית «Modal» — לא בשימוש אמיתי.", "להסתיר / למחוק לפני חנות", "", ""),
    ("X07", "תשתית", "טיפול בשגיאות רשת / 500 poll", S_PART, "P1 — חשוב", "poll הואט, prisma singleton. עדיין תלוי pooler 6543 ב-Render.", "", "השאר מעקב פתוח 2 דק", ""),
]

for i, row in enumerate(qa, 3):
    for c, v in enumerate(row, 1):
        ws.cell(i, c, v)
    ws.cell(i, 9, "")  # device

last = 2 + len(qa)
style_body(ws, 3, last, 10, status_col=4, pri_col=5)
add_status_dv(ws, "D", 3, last)
add_pri_dv(ws, "E", 3, last)
add_yn_dv(ws, "I", 3, last)
autosize(ws, [8, 16, 36, 32, 18, 62, 28, 32, 16, 22])
ws.freeze_panes = "A3"
ws.auto_filter.ref = f"A2:J{last}"
ws.row_dimensions[2].height = 24
for r in range(3, last + 1):
    ws.row_dimensions[r].height = 52

# ═══════════════════════════════════════════════════════════════════════════
# 03 TEXT FIELDS
# ═══════════════════════════════════════════════════════════════════════════
ws = wb.create_sheet("03_שדות_טקסט")
ws.sheet_view.rightToLeft = True
ws.merge_cells("A1:H1")
ws["A1"] = "כל תיבות הטקסט באפליקציה"
ws["A1"].font = font_title
ws["A1"].fill = fills["title"]
th = ["ID", "מסך", "שדה", "חובה?", "ולידציה בקוד", "סטטוס", "UX מקלדת / יישור", "נבדק במכשיר"]
for i, h in enumerate(th, 1):
    ws.cell(2, i, h)
style_header(ws, 2, 8)

fields = [
    ("F01", "התחברות", "אימייל", "כן (אם טופס פתוח)", "לא ריק; סוג email", S_CODE, "LTR מומלץ — בדוק במקלדת עברית", ""),
    ("F02", "התחברות", "סיסמה", "כן", "לא ריק, secure", S_CODE, "secureTextEntry", ""),
    ("F03", "הרשמה", "אימייל / סיסמה / אימות", "כן", "קיים בטופס", S_CODE, "כמו התחברות", ""),
    ("F04", "תיקון ש3", "שם לקוח", "כן", "trim + שגיאה אדומה", S_CODE, "יישור ימין", ""),
    ("F05", "תיקון ש3", "טלפון", "כן", "ספרות, אורך", S_CODE, "number-pad", ""),
    ("F06", "תיקון ש3", "חיפוש עיר", "כן (בחירה)", "רשימה סגורה", S_CODE, "modal חיפוש", ""),
    ("F07", "תיקון ש3", "חיפוש / הקלדת רחוב", "כן", "בחירה או ידני", S_PART, "תלוי API רחובות", ""),
    ("F08", "תיקון ש3", "מספר בית", "כן", "ניקוי תווים, max 8", S_CODE, "number-pad", ""),
    ("F09", "תיקון ש3", "אימייל אופציונלי", "לא", "regex אם מלא", S_CODE, "", ""),
    ("F10", "תיקון ש3", "תיאור תקלה חופשי", "לא", "multiline", S_CODE, "minHeight 80, ימין", ""),
    ("F11", "עריכת פרופיל", "שם", "כן", "trim", S_CODE, "ימין", ""),
    ("F12", "עריכת פרופיל", "טלפון", "כן", "ספרות בלבד עד 10", S_CODE, "phone-pad", ""),
    ("F13", "עריכת פרופיל", "אימייל", "לא ניתן לעריכה", "disabled", S_OK, "אפור", ""),
    ("F14", "פרופיל טכנאי", "ביו", "לא", "max 200", S_CODE, "multiline", ""),
    ("F15", "סיום עבודה טכנאי", "מחיר סופי", "כן לסיום", "מספר", S_CODE, "מקלדת מספרים", ""),
    ("F16", "סיום עבודה טכנאי", "שם חלק + מחיר חלק", "לא", "רשימה דינמית", S_CODE, "", ""),
    ("F17", "סיום עבודה טכנאי", "הערות", "לא", "טקסט", S_CODE, "", ""),
    ("F18", "תשלום נוסף", "תיאור תיקון נוסף", "כן", "min 2 בשרת", S_CODE, "multiline ימין", ""),
    ("F19", "תשלום נוסף", "סכום ₪", "כן", "10–50000 בשרת", S_CODE, "number-pad", ""),
    ("F20", "דירוג לקוח", "משוב חופשי", "לא", "multiline", S_CODE, "job-complete", ""),
    ("F21", "submit-review", "טקסט ביקורת", "לא", "קיים", S_CODE, "מסך נפרד", ""),
    ("F22", "צ'אט", "הודעה", "כן לשליחה", "לא ריק", S_CODE, "KeyboardAvoiding", ""),
    ("F23", "משיכת כסף", "סכום + בנק + סניף + חשבון + שם", "כן", "סכימה zod בשרת (סניף 3 ספרות וכו')", S_CODE, "שדות רבים — בדוק שגיאות עברית", ""),
    ("F24", "גלובלי", "כפתור סיום מקלדת (iOS)", "—", "KeyboardDoneToolbar", S_CODE, "«סיום» סוגר מקלדת", ""),
]
for i, row in enumerate(fields, 3):
    for c, v in enumerate(row, 1):
        ws.cell(i, c, v)
lf = 2 + len(fields)
style_body(ws, 3, lf, 8, status_col=6)
add_status_dv(ws, "F", 3, lf)
add_yn_dv(ws, "H", 3, lf)
autosize(ws, [8, 20, 28, 18, 28, 32, 28, 16])
ws.freeze_panes = "A3"
ws.auto_filter.ref = f"A2:H{lf}"

# ═══════════════════════════════════════════════════════════════════════════
# 04 UX
# ═══════════════════════════════════════════════════════════════════════════
ws = wb.create_sheet("04_UX_UI")
ws.sheet_view.rightToLeft = True
ws.merge_cells("A1:F1")
ws["A1"] = "סקירת UX / UI — מה מלוטש ומה נשאר מיושן"
ws["A1"].font = font_title
ws["A1"].fill = fills["title"]
for i, h in enumerate(["מסך / אזור", "רמת גימור", "מה טוב", "מה צורם", "עדיפות", "המלצה"], 1):
    ws.cell(2, i, h)
style_header(ws, 2, 6)

ux = [
    ("Splash + Loading + Sign-in", "גבוה", "אותה שפה: אופנוע ניאון, לוגו, ירוק", "Splash נייטיבי דורש rebuild", "P2 — שיפור", "ודא בבילד 1.8.1"),
    ("בית לקוח", "גבוה", "הדר זכוכית, CTA זוהר, מפה מלאה", "כרטיס הזמנה פעילה עדיין «כרטיס לבן»", "P2 — שיפור", "ליישר Hero לזכוכית אם רוצים אחידות"),
    ("דיווח תקלה — מעטפת", "גבוה", "הדר+פרוגרס+CTA זכוכית", "שדות טופס כתובת עדיין אפורי-מערכת", "P2 — שיפור", "אופציונלי"),
    ("בחירת טכנאי", "גבוה", "הדר, כרטיסים, כפתור pill", "מיון/ריק — טוב", "P2 — שיפור", ""),
    ("אישור הזמנה / ConfirmModal", "גבוה", "זכוכית כחולה, טקסט גדול", "כל המודאלים באותו סגנון עכשיו", "P2 — שיפור", ""),
    ("המתנה / תשלום / סיום תיקון", "גבוה", "מסכי המרה כהים + CTA גדול", "ניסיונות «ריבוע באמצע» בוטלו — נשאר מלבן גדול למטה", "P1 — חשוב", "אל תחזור לריבוע ממורכז"),
    ("מעקב אחרי תשלום (מפה)", "בינוני", "כפתורי תחתית הוגדלו", "השטח הלבן של ה-sheet מול מפת כהה — שתי שפות", "P2 — שיפור", ""),
    ("עבודה פעילה טכנאי", "בינוני-גבוה", "כהה, טיימליין ברור, תיאור תקלה חדש", "לא זכוכית כמו הלקוח", "P2 — שיפור", "אפשר להשאיר — שפה «עבודה»"),
    ("דשבורד טכנאי / רווחים / הזמנות / אדמין", "נמוך", "פונקציונלי", "נראה דור קודם (לבן/כחול גנרי)", "P1 — חשוב", "סבב עיצוב הבא"),
    ("פרופיל לקוח + משפטי", "נמוך-בינוני", "כל הלינקים המשפטיים קיימים", "רשימות הגדרות גנריות", "P2 — שיפור", ""),
    ("צ'אט / ביקורות / משיכה", "נמוך", "עובד בקוד", "לא מותג", "P2 — שיפור", ""),
    ("עברית / RTL", "בינוני", "ברירת מחדל עברית", "flex-end/row לפעמים מתהפך; מסכים מעורבים", "P1 — חשוב", "בדוק כל מסך חדש ב-RTL"),
    ("נגישות", "חלקי", "חלק מהכפתורים עם accessibilityLabel", "לא שיטתי, אין Dynamic Type", "P2 — שיפור", "לפני סקירת Apple — לפחות לכפתורים ראשיים"),
    ("מצב כהה מערכתי", "חלקי", "toggle בפרופיל", "מסכים רבים צבע קבוע", "P2 — שיפור", "או לכבות את ה-toggle עד שיהיה אמיתי"),
    ("עקביות כפתורים", "חלקי", "CTA לקוח חזקים", "גדלים שונים בין מסכים (בכוונה חלקית)", "P2 — שיפור", "תעד כלל: ראשי ≥56, המרה ≥72"),
]
for i, row in enumerate(ux, 3):
    for c, v in enumerate(row, 1):
        ws.cell(i, c, v)
style_body(ws, 3, 2 + len(ux), 6, pri_col=5)
autosize(ws, [32, 16, 40, 48, 18, 36])
ws.freeze_panes = "A3"
for r in range(3, 3 + len(ux)):
    ws.row_dimensions[r].height = 40

# ═══════════════════════════════════════════════════════════════════════════
# 05 APPLE 100
# ═══════════════════════════════════════════════════════════════════════════
ws = wb.create_sheet("05_אפל_100")
ws.sheet_view.rightToLeft = True
ws.merge_cells("A1:G1")
ws["A1"] = "App Store — 100% מוכנות  |  אפליקציה 6775995887  |  com.ebikeland.app"
ws["A1"].font = font_title
ws["A1"].fill = fills["title"]
ah = ["ID", "פריט", "חובה?", "סטטוס", "פירוט / איך", "נעשה?", "הערה"]
for i, h in enumerate(ah, 1):
    ws.cell(2, i, h)
style_header(ws, 2, 7)

apple = [
    ("AS01", "חשבון Apple Developer פעיל + הסכמים חתומים", "כן", S_PART, "בעבר submit נחסם על הסכם. בדוק Agreements, Tax, Banking ב-ASC.", "", "בלי זה אין TF/Review"),
    ("AS02", "אפליקציה ב-ASC עם bundle נכון", "כן", S_OK, "אפליקציה חדשה 6775995887 = com.ebikeland.app (לא הישנה 6759334424)", "", "eas.json ascAppId תואם"),
    ("AS03", "גרסת קוד בחנות (marketing + build)", "כן", S_CODE, "1.8.1 (24). ודא שבילד EAS הסתיים והופיע ב-TestFlight.", "", "לינק: expo.dev .../builds/6931d93c-…"),
    ("AS04", "TestFlight פנימי — אתה + אלביס", "כן", S_PART, "התקן 1.8.1, עבור E2E מלא.", "", ""),
    ("AS05", "אייקון 1024, בלי שקיפות, בלי rounded יתר", "כן", S_CODE, "icon.png 1024. בדוק ב-TF איך נראה על השארית.", "", ""),
    ("AS06", "Splash תקין על iPhone אמיתי", "כן", S_CODE, "תמונת cover כהה.", "", ""),
    ("AS07", "צילומי מסך 6.7\" ו-6.5\" (לפחות)", "כן", S_MISSING, "צלם מהבילד החדש: בית, תיקון, תשלום, מעקב, טכנאי.", "", "בלי זה אין Submit"),
    ("AS08", "תיאור, כותרת משנה, מילות מפתח, קטגוריה", "כן", S_MISSING, "לכתוב בעברית. קטגוריה Lifestyle / Travel?", "", ""),
    ("AS09", "Privacy Policy URL ציבורי (https)", "כן", S_MISSING, "יש טקסט בתוך האפליקציה בלבד. צריך דף ב-ebikeland.com", "", "בלוקר סקירה"),
    ("AS10", "Support URL", "כן", S_MISSING, "למשל ebikeland.com/support או וואטסאפ עמוד", "", ""),
    ("AS11", "Privacy Nutrition Labels ב-ASC", "כן", S_MISSING, "לסמן: Location, Contact Info, Photos, Purchases, Identifiers, Diagnostics.", "", "תואם למדיניות באפליקציה"),
    ("AS12", "Sign in with Apple (אם יש Google)", "כן אם Google פעיל", S_PART, "כפתור קיים. חייב לעבוד במכשיר + callback ב-Render.", "", "Guideline 4.8"),
    ("AS13", "מחיקת חשבון בתוך האפליקציה", "כן", S_CODE, "קיים בפרופיל + DELETE /api/users/me", "", "Guideline 5.1.1"),
    ("AS14", "הסברי הרשאות (מצלמה, תמונות, מיקום תמיד)", "כן", S_OK, "infoPlist בעברית/אנגלית תקינים", "", ""),
    ("AS15", "Background location — הצדקה אמיתית", "כן", S_PART, "רק לטכנאי זמין. הסבר קיים. Apple עלולה לשאול.", "", "הכן תשובה ל-Review Notes"),
    ("AS16", "רכישות / תשלום מחוץ ל-IAP", "כן (שירות בעולם האמיתי)", S_PART, "שירות תיקון פיזי → מותר PayMe. להסביר ב-Review Notes.", "", "Guideline 3.1.3(d) / 3.1.5"),
    ("AS17", "אין כפתור «תשלום דמו» בבילד חנות", "כן", S_OK, "eas production: MOCK + SIMULATE = false בצד הלקוח", "", "אבל השרת עדיין mock — ראה בלוקרים"),
    ("AS18", "Export compliance / הצפנה", "כן", S_OK, "ITSAppUsesNonExemptEncryption=false", "", ""),
    ("AS19", "גיל / אין ילדים", "כן", S_MISSING, "4+ או 12+ לפי תוכן. אין UGC מסוכן.", "", ""),
    ("AS20", "Review Notes + חשבון דמו לסוקר", "כן", S_MISSING, "תן a@abc.com + tech + הסבר זרימה דו-משתמשת", "", "בלי זה הסוקר נתקע"),
    ("AS21", "וידאו תצוגה מקדימה (אופציונלי)", "לא", S_MISSING, "מומלץ", "", ""),
    ("AS22", "Copyright / שם חברה", "כן", S_PART, "eBike Land / ebikeland.com במסמכים", "", "יישור מול דף ASC"),
    ("AS23", "Push — הרשאה לא חובה לכניסה", "כן", S_CODE, "לא חוסם login", "", ""),
    ("AS24", "אין קריסה בפתיחה / רשימת טכנאים", "כן", S_PART, "תוקנו קריסות distance/location. חובה בדיקת רגרסיה ב-TF.", "", ""),
    ("AS25", "Submit for Review + תשובות לשאלות", "כן", S_MISSING, "אחרי Internal TF יציב", "", ""),
]
for i, row in enumerate(apple, 3):
    for c, v in enumerate(row, 1):
        ws.cell(i, c, v)
la = 2 + len(apple)
style_body(ws, 3, la, 7, status_col=4)
add_status_dv(ws, "D", 3, la)
add_yn_dv(ws, "F", 3, la)
autosize(ws, [8, 42, 16, 32, 62, 12, 28])
ws.freeze_panes = "A3"
ws.auto_filter.ref = f"A2:G{la}"
for r in range(3, la + 1):
    ws.row_dimensions[r].height = 38

# ═══════════════════════════════════════════════════════════════════════════
# 06 GOOGLE 100
# ═══════════════════════════════════════════════════════════════════════════
ws = wb.create_sheet("06_גוגל_100")
ws.sheet_view.rightToLeft = True
ws.merge_cells("A1:G1")
ws["A1"] = "Google Play — 100% מוכנות  |  package com.ebike.app  |  עדיין אין אפליקציה בקונסול"
ws["A1"].font = font_title
ws["A1"].fill = fills["title"]
for i, h in enumerate(ah, 1):
    ws.cell(2, i, h)
style_header(ws, 2, 7)

google = [
    ("GP01", "חשבון Play Console + תשלום $25", "כן", S_MISSING, "play.google.com/console — הצעד הראשון", "", "יכול לקחת עד 48 שעות"),
    ("GP02", "פרופיל מפתח / ארגון / D-U-N-S אם חברה", "כן", S_MISSING, "יחיד או ארגון", "", ""),
    ("GP03", "Create app: eBike, he, Free, App", "כן", S_MISSING, "package חייב com.ebike.app — נעול אחרי הפרסום הראשון", "", "שונה מ-iOS bundle בכוונה"),
    ("GP04", "AAB production מ-EAS", "כן", S_CODE, "eas build --platform android --profile production", "", "רק אחרי ש-GP03 קיים"),
    ("GP05", "חתימה / Play App Signing", "כן", S_CODE, "EAS מנהל keystore. שמור credentials.", "", "אל תאבד את המפתח"),
    ("GP06", "העלאה ל-Internal testing (draft)", "כן", S_CODE, "eas.json כבר track=internal, releaseStatus=draft", "", "מתאים לבדיקה ראשונה"),
    ("GP07", "Service account ל-EAS Submit (אופציונלי)", "לא בהתחלה", S_MISSING, "אפשר גם העלאת AAB ידנית", "", ""),
    ("GP08", "Store listing: כותרת, תיאור קצר (80), ארוך", "כן", S_MISSING, "עברית", "", ""),
    ("GP09", "אייקון 512×512 + feature graphic 1024×500", "כן", S_MISSING, "adaptiveIcon לבן + icon.png — צריך נכסי חנות", "", ""),
    ("GP10", "צילומי מסך טלפון (לפחות 2, עדיף 8)", "כן", S_MISSING, "מאותו ביילד כמו iOS", "", ""),
    ("GP11", "Privacy Policy URL", "כן", S_MISSING, "אותו דף כמו אפל", "", "בלוקר"),
    ("GP12", "Data safety form", "כן", S_MISSING, "מיקום מדויק, תמונות, מידע אישי, תשלומים, מזהי מכשיר", "", ""),
    ("GP13", "Ads declaration = No", "כן", S_MISSING, "אין פרסומות", "", ""),
    ("GP14", "Target audience / Kids = לא לילדים", "כן", S_MISSING, "", "", ""),
    ("GP15", "News / COVID / COVID app = לא", "כן", S_MISSING, "", "", ""),
    ("GP16", "Content rating (IARC)", "כן", S_MISSING, "שאלון קצר בקונסול", "", ""),
    ("GP17", "Foreground + background location הצהרה", "כן", S_PART, "רק טכנאי זמין. למלא טופס הרשאות רגישות.", "", "Google מחמירה כמו אפל"),
    ("GP18", "Photo/Camera policy", "כן", S_PART, "צילום תקלה אופציונלי", "", ""),
    ("GP19", "פיננסים / תשלומים מחוץ ל-Play", "כן", S_PART, "שירות פיזי → PayMe מותר. להצהיר.", "", ""),
    ("GP20", "חשבון בדיקה לסוקר + הוראות דו-תפקיד", "כן", S_MISSING, "כמו אפל", "", ""),
    ("GP21", "מחיקת חשבון / בקשת מחיקה", "כן", S_CODE, "בתוך האפליקציה", "", "Play דורש גם דרך ברורה"),
    ("GP22", "בדיקת Pre-launch report (מעבדת Firebase)", "מומלץ", S_MISSING, "אחרי העלאת AAB", "", "לתקן קריסות Android"),
    ("GP23", "Android  permissions ישנות (storage)", "לבדוק", S_PART, "READ/WRITE_EXTERNAL_STORAGE עלול להידחות ב-API חדש", "", "לנקות לפני target חדש"),
    ("GP24", "מפות Google — SHA-1 restriction", "כן", S_PART, "להוסיף SHA של App Signing ב-Cloud Console", "", "בלי זה מפה שחורה בפרודקשן"),
    ("GP25", "קידום ל-Closed/Open/Production", "כן", S_MISSING, "רק אחרי Internal יציב + PayMe אמיתי", "", "אל תצא Production על mock"),
]
for i, row in enumerate(google, 3):
    for c, v in enumerate(row, 1):
        ws.cell(i, c, v)
lg = 2 + len(google)
style_body(ws, 3, lg, 7, status_col=4)
add_status_dv(ws, "D", 3, lg)
add_yn_dv(ws, "F", 3, lg)
autosize(ws, [8, 46, 16, 32, 58, 12, 28])
ws.freeze_panes = "A3"
ws.auto_filter.ref = f"A2:G{lg}"
for r in range(3, lg + 1):
    ws.row_dimensions[r].height = 38

# ═══════════════════════════════════════════════════════════════════════════
# 07 BLOCKERS
# ═══════════════════════════════════════════════════════════════════════════
ws = wb.create_sheet("07_בלוקרים")
ws.sheet_view.rightToLeft = True
ws.merge_cells("A1:F1")
ws["A1"] = "רק זה חוסם אותך — לפי סדר ביצוע"
ws["A1"].font = font_title
ws["A1"].fill = fills["title"]
for i, h in enumerate(["#", "בלוקר", "למה זה חוסם", "מה לעשות", "שייך ל", "בוצע?"], 1):
    ws.cell(2, i, h)
style_header(ws, 2, 6)

blocks = [
    (1, "Render עדיין MOCK_PAYMENTS=true", "בחנות הלקוח «משלם» בדף מזויף. זה לא חוקי ולא אמיתי.", "Render env: MOCK_PAYMENTS=false, מפתחות PayMe חיים/סנדבוקס לפי החלטה, BACKEND_URL ציבורי, DB pooler 6543. Deploy. אמת /health.", "שרת", ""),
    (2, "E2E מלא לא רץ כבר ~חודש", "אי אפשר לדעת אם 1.8.1 באמת יציב במכשיר.", "TestFlight 1.8.1: לקוח+טכנאי מקצה לקצה כולל תשלום.", "אתה + אלביס", ""),
    (3, "אין דף מדיניות פרטיות ב-HTTPS", "אפל וגוגל דוחים בלי URL.", "העלה את legal-content ל-ebikeland.com/privacy (ו-/terms /cancellation).", "אתר", ""),
    (4, "צילומי מסך + טקסט חנות", "אין Submit בלי נכסים.", "צלם מ-TF אחרי שהזרימה עובדת.", "שיווק", ""),
    (5, "Google / Apple login במכשיר אמיתי", "אם הכפתורים גלויים — חייבים לעבוד. סימולטור לא מספיק.", "בדוק OAuth ב-Render + callback URLs.", "שרת + מכשיר", ""),
    (6, "טכנאי מופיע אצל הלקוח", "בלי זה אין מוצר.", "זמין+מיקום+אישור אדמין. רדיוס. אל תסגור את האפליקציה בלי heartbeat.", "E2E", ""),
    (7, "Play Console לא קיים", "אין לאן להעלות Android.", "חשבון $25 → Create app com.ebike.app.", "גוגל", ""),
    (8, "הסכמי Apple / מס / בנק", "חוסם העלאה כמו שקרה בעבר.", "ASC → Business + Agreements.", "אפל", ""),
    (9, "חשבון דמו לסוקר + הוראות דו-משתמש", "הסוקר לא יכול לקבל הזמנה בלי טכנאי שני.", "כתוב צעד-אחר-צעד בעברית+אנגלית.", "מסמך", ""),
    (10, "Commit של שינויי UI/extra-pay לבילד", "אם EAS לא ארז את הקוד המקומי — TF ישן.", "git status; אם dirty — commit ואז ביילד חדש.", "גיט", ""),
]
for i, row in enumerate(blocks, 3):
    for c, v in enumerate(row, 1):
        ws.cell(i, c, v)
        ws.cell(i, c).fill = fills["p0"]
style_body(ws, 3, 2 + len(blocks), 6)
add_yn_dv(ws, "F", 3, 2 + len(blocks))
autosize(ws, [6, 40, 42, 62, 16, 12])
ws.freeze_panes = "A3"
for r in range(3, 3 + len(blocks)):
    ws.row_dimensions[r].height = 48

# ═══════════════════════════════════════════════════════════════════════════
# 08 SUMMARY (formulas)
# ═══════════════════════════════════════════════════════════════════════════
ws = wb.create_sheet("סיכום", 1)
ws.sheet_view.rightToLeft = True
ws.merge_cells("A1:C1")
ws["A1"] = "סיכום חי — מתעדכן כשמשנים סטטוסים בגיליון המאסטר"
ws["A1"].font = font_title
ws["A1"].fill = fills["title"]
ws.row_dimensions[1].height = 28

ws["A3"] = "מדד"
ws["B3"] = "מספר"
style_header(ws, 3, 2)

metrics = [
    ("סה״כ פריטי QA (מאסטר)", f"=COUNTA('02_מאסטר_QA'!A3:A{last})"),
    ("עובד (ידוע)", f"=COUNTIF('02_מאסטר_QA'!D:D,\"{S_OK}\")"),
    ("קיים בקוד — ממתין למכשיר", f"=COUNTIF('02_מאסטר_QA'!D:D,\"{S_CODE}\")"),
    ("חלקי", f"=COUNTIF('02_מאסטר_QA'!D:D,\"{S_PART}\")"),
    ("לא עובד", f"=COUNTIF('02_מאסטר_QA'!D:D,\"{S_BROKEN}\")"),
    ("בלוקר פרודקשן", f"=COUNTIF('02_מאסטר_QA'!D:D,\"{S_BLOCK}\")"),
    ("חסר לחנות", f"=COUNTIF('02_מאסטר_QA'!D:D,\"{S_MISSING}\")"),
    ("UX חלקי", f"=COUNTIF('02_מאסטר_QA'!D:D,\"{S_UX}\")"),
    ("נבדק במכשיר = כן", f"=COUNTIF('02_מאסטר_QA'!I:I,\"כן\")"),
    ("פריטי אפל", f"=COUNTA('05_אפל_100'!A3:A{la})"),
    ("אפל — נעשה כן", f"=COUNTIF('05_אפל_100'!F:F,\"כן\")"),
    ("פריטי גוגל", f"=COUNTA('06_גוגל_100'!A3:A{lg})"),
    ("גוגל — נעשה כן", f"=COUNTIF('06_גוגל_100'!F:F,\"כן\")"),
]
for i, (name, formula) in enumerate(metrics, 4):
    ws.cell(i, 1, name).font = font_b
    ws.cell(i, 2, formula).font = Font(name="Calibri", size=14, bold=True)
    for c in range(1, 3):
        ws.cell(i, c).border = thin
        ws.cell(i, c).alignment = wrap

ws["A18"] = "אחוז מאסטר שסומן «עובד»"
ws["B18"] = f"=IF(B4=0,0,B5/B4)"
ws["B18"].number_format = "0%"
ws["A19"] = "אחוז אפל שסומן נעשה"
ws["B19"] = f"=IF(B13=0,0,B14/B13)"
ws["B19"].number_format = "0%"
ws["A20"] = "אחוז גוגל שסומן נעשה"
ws["B20"] = f"=IF(B15=0,0,B16/B15)"
ws["B20"].number_format = "0%"
for r in range(18, 21):
    ws.cell(r, 1).font = font_b
    ws.cell(r, 1).fill = fills["note"]
    ws.cell(r, 2).fill = fills["note"]
    ws.cell(r, 1).border = thin
    ws.cell(r, 2).border = thin

ws.merge_cells("A22:B26")
ws["A22"] = (
    "קריאת מצב נוכחית (אוגוסט 2026):\n"
    "• המוצר כמעט שלם בקוד — הזרימה הראשית קיימת.\n"
    "• פרודקשן עדיין בגובה mock — אסור לצאת לחנות ככה.\n"
    "• UI של הלקוח בזרימה הראשית שודרג; פרופיל/טכנאי/אדמין מאחור.\n"
    "• הצעד הבא שלך: TestFlight 1.8.1 + הדלקת PayMe ב-Render + סימון העמודה «נבדק במכשיר»."
)
ws["A22"].alignment = wrap
ws["A22"].fill = fills["note"]
ws.row_dimensions[22].height = 80

autosize(ws, [42, 18])

# ═══════════════════════════════════════════════════════════════════════════
# print / freeze
# ═══════════════════════════════════════════════════════════════════════════
for s in wb.worksheets:
    s.page_setup.orientation = "landscape"
    s.page_setup.fitToPage = True
    s.page_setup.fitToWidth = 1
    s.page_setup.fitToHeight = 0
    s.sheet_properties.pageSetUpPr.fitToPage = True
    s.page_setup.paperSize = s.PAPERSIZE_A4
    s.print_options.horizontalCentered = True
    s.oddHeader.left.text = "eBike QA"
    s.oddFooter.right.text = "עמוד &P מתוך &N"

wb.save(OUT)
print("Wrote", OUT)
print("QA rows", len(qa), "fields", len(fields), "apple", len(apple), "google", len(google))
